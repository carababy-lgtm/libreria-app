#!/usr/bin/env python3
import os, io, csv, threading
from flask import (Flask, render_template, request, redirect,
                   url_for, session, jsonify, send_file, flash)

from config   import PASSWORD, STANZE, RIPIANI, DB_PATH
from database import (init_db, db_all, db_search, db_get,
                      db_insert, db_update, db_delete, db_stats, gia_presente)
from lookup   import lookup_espanso, lookup_singolo
from backup   import esegui_backup, ripristina_da_github, stato_backup, avvia_backup_automatico

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "libreria-secret-2024")


# ══════════════════════════════════════════════════════════
#  AUTH
# ══════════════════════════════════════════════════════════
def logged_in():
    return session.get("auth") is True


@app.route("/login", methods=["GET", "POST"])
def login():
    errore = ""
    if request.method == "POST":
        if request.form.get("password") == PASSWORD:
            session.clear()
            session["auth"] = True
            return redirect(url_for("catalogo"))
        errore = "Password errata."
    return render_template("login.html", errore=errore)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ══════════════════════════════════════════════════════════
#  CATALOGO
# ══════════════════════════════════════════════════════════
@app.route("/")
def index():
    return redirect(url_for("catalogo"))


@app.route("/catalogo")
def catalogo():
    if not logged_in():
        return redirect(url_for("login"))
    q      = request.args.get("q", "").strip()
    stanza = request.args.get("stanza", "")
    rows   = db_search(q, stanza) if (q or stanza) else db_all()
    stats, dist_stanze = db_stats()
    return render_template("catalogo.html",
                           libri=rows, q=q, stanza=stanza,
                           stanze=STANZE, stats=stats,
                           dist_stanze=dist_stanze,
                           backup_stato=stato_backup())


# ══════════════════════════════════════════════════════════
#  DETTAGLIO
# ══════════════════════════════════════════════════════════
@app.route("/libro/<int:book_id>")
def dettaglio(book_id):
    if not logged_in():
        return redirect(url_for("login"))
    libro = db_get(book_id)
    if not libro:
        flash("Libro non trovato.", "warning")
        return redirect(url_for("catalogo"))
    return render_template("dettaglio.html", libro=libro)


# ══════════════════════════════════════════════════════════
#  AGGIUNGI / MODIFICA
# ══════════════════════════════════════════════════════════
@app.route("/nuovo", methods=["GET", "POST"])
def nuovo():
    if not logged_in():
        return redirect(url_for("login"))
    if request.method == "POST":
        titolo = request.form.get("titolo", "").strip()
        if not titolo:
            flash("Il titolo è obbligatorio.", "danger")
            return render_template("form_libro.html",
                                   libro={}, stanze=STANZE, ripiani=RIPIANI,
                                   titolo_pag="Nuovo libro")
        db_insert(_form_to_dict(request.form))
        flash(f"«{titolo}» aggiunto.", "success")
        return redirect(url_for("catalogo"))
    # GET — campi sempre vuoti
    return render_template("form_libro.html",
                           libro={}, stanze=STANZE, ripiani=RIPIANI,
                           titolo_pag="Nuovo libro")


@app.route("/modifica/<int:book_id>", methods=["GET", "POST"])
def modifica(book_id):
    if not logged_in():
        return redirect(url_for("login"))
    libro = db_get(book_id)
    if not libro:
        flash("Libro non trovato.", "warning")
        return redirect(url_for("catalogo"))
    if request.method == "POST":
        titolo = request.form.get("titolo", "").strip()
        if not titolo:
            flash("Il titolo è obbligatorio.", "danger")
            return render_template("form_libro.html",
                                   libro=libro, stanze=STANZE, ripiani=RIPIANI,
                                   titolo_pag=f"Modifica [{book_id}]")
        db_update(book_id, _form_to_dict(request.form))
        flash("Modifiche salvate.", "success")
        return redirect(url_for("dettaglio", book_id=book_id))
    return render_template("form_libro.html",
                           libro=libro, stanze=STANZE, ripiani=RIPIANI,
                           titolo_pag=f"Modifica [{book_id}]")


@app.route("/elimina/<int:book_id>", methods=["POST"])
def elimina(book_id):
    if not logged_in():
        return redirect(url_for("login"))
    libro = db_get(book_id)
    db_delete(book_id)
    flash(f"«{libro.get('titolo','')}» eliminato.", "info")
    return redirect(url_for("catalogo"))


def _form_to_dict(form):
    return {
        "titolo":      form.get("titolo", "").strip(),
        "autore":      form.get("autore", "").strip(),
        "argomento":   form.get("argomento", "").strip(),
        "anno":        form.get("anno", "").strip(),
        "editore":     form.get("editore", "").strip(),
        "isbn":        form.get("isbn", "").strip(),
        "descrizione": form.get("descrizione", "").strip(),
        "stanza":      form.get("stanza", "").strip(),
        "ripiano":     form.get("ripiano", "").strip(),
        "note":        form.get("note", "").strip(),
    }


# ══════════════════════════════════════════════════════════
#  LOOKUP ONLINE (AJAX)
# ══════════════════════════════════════════════════════════
@app.route("/api/lookup")
def api_lookup():
    if not logged_in():
        return jsonify({"error": "Non autorizzato"}), 401
    q = request.args.get("q", "").strip()
    if not q:
        return jsonify({"error": "Query vuota"}), 400
    result, source = lookup_singolo(q)
    if result:
        return jsonify({"ok": True, "data": result, "source": source})
    return jsonify({"ok": False, "source": ""})


# ══════════════════════════════════════════════════════════
#  IMPORTAZIONE DA FILE EXCEL / CSV
# ══════════════════════════════════════════════════════════
@app.route("/importa", methods=["GET", "POST"])
def importa():
    if not logged_in():
        return redirect(url_for("login"))
    if request.method == "GET":
        return render_template("importa.html", stanze=STANZE, ripiani=RIPIANI,
                               risultati=None, elaborati=0)

    f = request.files.get("file")
    if not f or f.filename == "":
        flash("Seleziona un file CSV o Excel.", "warning")
        return render_template("importa.html", stanze=STANZE, ripiani=RIPIANI,
                               risultati=None, elaborati=0)

    stanza_def  = request.form.get("stanza", "").strip()
    ripiano_def = request.form.get("ripiano", "").strip()

    # Leggi righe dal file
    righe = _leggi_file(f)
    if righe is None:
        flash("Formato file non supportato. Usa CSV o Excel (.xlsx).", "danger")
        return render_template("importa.html", stanze=STANZE, ripiani=RIPIANI,
                               risultati=None, elaborati=0)

    risultati = []
    for riga in righe:
        titolo  = str(riga.get("titolo", "")).strip()
        if not titolo:
            continue
        autore  = str(riga.get("autore",  "")).strip()
        anno    = str(riga.get("anno",    "")).strip()
        editore = str(riga.get("editore", "")).strip()
        stanza  = str(riga.get("stanza",  "")).strip() or stanza_def
        ripiano = str(riga.get("ripiano", "")).strip() or ripiano_def

        if gia_presente(titolo):
            risultati.append({"titolo": titolo, "autore": autore,
                               "stato": "già presente", "classe": "warning"})
            continue

        # Lookup online
        extra, source = lookup_espanso(autore, titolo)

        data = {
            "titolo":      titolo,
            "autore":      autore or (extra.get("autore", "") if extra else ""),
            "argomento":   extra.get("argomento", "") if extra else "",
            "descrizione": extra.get("descrizione", "") if extra else "",
            "anno":        anno or (extra.get("anno", "") if extra else ""),
            "editore":     editore or (extra.get("editore", "") if extra else ""),
            "isbn":        extra.get("isbn", "") if extra else "",
            "stanza":      stanza,
            "ripiano":     ripiano,
            "note":        "",
        }
        db_insert(data)
        stato  = f"importato ({source})" if extra else "importato (senza dati online)"
        classe = "success" if extra else "warning"
        risultati.append({"titolo": titolo, "autore": data["autore"],
                           "stato": stato, "classe": classe})

    return render_template("importa.html", stanze=STANZE, ripiani=RIPIANI,
                           risultati=risultati, elaborati=len(risultati))


def _leggi_file(f):
    """Legge CSV o XLSX e restituisce lista di dict con chiavi minuscole."""
    nome = f.filename.lower()
    try:
        if nome.endswith(".csv"):
            contenuto = f.read().decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(contenuto))
            return [{k.strip().lower(): v for k, v in row.items()} for row in reader]
        elif nome.endswith(".xlsx") or nome.endswith(".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
            ws = wb.active
            intestazioni = [str(c.value).strip().lower() if c.value else "" 
                           for c in next(ws.iter_rows(min_row=1, max_row=1))]
            righe = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                d = {intestazioni[i]: (str(v).strip() if v is not None else "")
                     for i, v in enumerate(row)}
                righe.append(d)
            return righe
        else:
            return None
    except Exception:
        return None


# ══════════════════════════════════════════════════════════
#  SCARICA TEMPLATE CSV
# ══════════════════════════════════════════════════════════
@app.route("/template-csv")
def template_csv():
    if not logged_in():
        return redirect(url_for("login"))
    output = io.StringIO()
    w = csv.writer(output)
    w.writerow(["titolo", "autore", "anno", "editore", "stanza", "ripiano"])
    w.writerow(["L'alchimista", "Paulo Coelho", "1988", "Bompiani", "STUDIO", "A1"])
    w.writerow(["Gomorra", "Roberto Saviano", "", "", "STUDIO", "B2"])
    output.seek(0)
    return send_file(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        mimetype="text/csv",
        as_attachment=True,
        download_name="template_libri.csv"
    )


# ══════════════════════════════════════════════════════════
#  BACKUP
# ══════════════════════════════════════════════════════════
@app.route("/backup/esegui", methods=["POST"])
def backup_esegui():
    if not logged_in():
        return redirect(url_for("login"))
    ok, msg = esegui_backup()
    flash(msg, "success" if ok else "danger")
    return redirect(url_for("catalogo"))


@app.route("/backup/scarica")
def backup_scarica():
    if not logged_in():
        return redirect(url_for("login"))
    return send_file(DB_PATH, as_attachment=True, download_name="libri.db")


@app.route("/backup/ripristina", methods=["POST"])
def backup_ripristina():
    if not logged_in():
        return redirect(url_for("login"))
    f = request.files.get("db_file")
    if f and f.filename.endswith(".db"):
        import sqlite3, tempfile, os
        tmp = tempfile.mktemp(suffix=".db")
        f.save(tmp)
        try:
            conn = sqlite3.connect(tmp)
            conn.execute("SELECT COUNT(*) FROM libri")
            conn.close()
            os.replace(tmp, DB_PATH)
            flash("DB ripristinato dal file caricato.", "success")
        except Exception:
            if os.path.exists(tmp):
                os.unlink(tmp)
            flash("File DB non valido.", "danger")
    else:
        ok, msg = ripristina_da_github()
        flash(msg, "success" if ok else "danger")
    return redirect(url_for("catalogo"))


# ══════════════════════════════════════════════════════════
#  AVVIO
# ══════════════════════════════════════════════════════════
if __name__ == "__main__":
    init_db()
    avvia_backup_automatico()
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
